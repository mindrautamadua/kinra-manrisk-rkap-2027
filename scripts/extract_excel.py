import openpyxl, json, os, sys, re
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

SRC_KK, SRC_TAX, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
os.makedirs(os.path.join(OUT, 'sheets'), exist_ok=True)
wbf = openpyxl.load_workbook(SRC_KK)                 # formulas
wbv = openpyxl.load_workbook(SRC_KK, data_only=True)  # cached values

import datetime
EPOCH = datetime.datetime(1899, 12, 30)
def val(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return round((v - EPOCH).total_seconds() / 86400, 6)
    if isinstance(v, datetime.date): return (datetime.datetime(v.year, v.month, v.day) - EPOCH).days
    if hasattr(v, 'isoformat'): return v.isoformat()
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v

def s(v):
    v = val(v)
    if v is None: return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v

# ---------- taksonomi ----------
wt = openpyxl.load_workbook(SRC_TAX, data_only=True).active
tax = []
for r in wt.iter_rows(min_row=2, values_only=True):
    if not r[6]: continue
    tax.append(dict(kode=s(r[6]), kategori=s(r[0]), def_kategori=s(r[1]), aspek=s(r[2]), def_aspek=s(r[3]),
                    risiko_korporat=s(r[4]), def_risiko_korporat=s(r[5]), risiko_spesifik=s(r[7]),
                    def_risiko_spesifik=s(r[8]), jumlah_unit=int(r[9] or 0)))
json.dump(tax, open(os.path.join(OUT, 'taksonomi.json'), 'w'), ensure_ascii=False)

# ---------- structured sheets ----------
V = lambda sh, a: s(wbv[sh][a].value)
sampul = dict(nama_entitas=V('Sampul','E10'), posisi_data=None, tanggal_penyusunan=None, versi=str(V('Sampul','E13') or '1'),
              catatan=V('Sampul','B15'))

ws = wbv['RAS']
ras = []
cur = None
for r in range(8, 18):
    no, kat = s(ws.cell(r,2).value), s(ws.cell(r,3).value)
    if kat: cur = (no, kat)
    if not s(ws.cell(r,4).value): continue
    ras.append(dict(no=int(cur[0]), kategori=cur[1], sikap_induk=s(ws.cell(r,4).value), pernyataan_induk=s(ws.cell(r,5).value),
                    sikap=s(ws.cell(r,6).value), pernyataan=s(ws.cell(r,7).value), rasionalisasi=s(ws.cell(r,8).value)))
ras_tambahan = []
for r in range(23, 33):
    if s(ws.cell(r,3).value):
        ras_tambahan.append(dict(kategori=s(ws.cell(r,3).value), sikap=s(ws.cell(r,4).value), pernyataan=s(ws.cell(r,5).value), rasionalisasi=s(ws.cell(r,6).value)))

ws = wbv['Kriteria Dampak']
kategori = []
for r in range(8, 16):
    full = s(ws.cell(r,2).value)
    m = re.match(r'(.*?)\s*\((.*)\)$', full)
    kategori.append(dict(no=int(ws.cell(r,3).value), nama=m.group(1).split('/')[0], nama_en=m.group(1) + ' (' + m.group(2) + ')', definisi=s(ws.cell(r,4).value)))
kriteria = []
cur = None
for r in range(20, 57):
    if s(ws.cell(r,2).value): cur = s(ws.cell(r,2).value)
    kriteria.append(dict(kategori=cur, kode=s(ws.cell(r,3).value), kriteria=s(ws.cell(r,4).value), perhitungan=s(ws.cell(r,5).value),
                         sangat_rendah=s(ws.cell(r,6).value), rendah=s(ws.cell(r,7).value), moderat=s(ws.cell(r,8).value),
                         tinggi=s(ws.cell(r,9).value), sangat_tinggi=s(ws.cell(r,10).value)))

ws = wbv['Risk Driver']
drivers = []
cur = None
for r in range(8, 106):
    if s(ws.cell(r,3).value): cur = s(ws.cell(r,3).value)
    if s(ws.cell(r,4).value):
        drivers.append(dict(kategori=cur, driver=s(ws.cell(r,4).value), bentuk=s(ws.cell(r,5).value), tambahan=False))
for r in range(111, 131):
    if s(ws.cell(r,3).value) and s(ws.cell(r,4).value):
        drivers.append(dict(kategori=s(ws.cell(r,3).value), driver=s(ws.cell(r,4).value), bentuk=None, tambahan=True))

ws = wbv['Risk Register']
register = []
for r in range(9, 209):
    row = [s(ws.cell(r,c).value) for c in range(3, 18)]
    if not any(row): continue
    keys = ['kategori','risk_driver','risk_event','deskripsi','root_cause','probabilitas','deskripsi_dampak','nilai_dampak',
            'kriteria_dampak','kategori_dampak','tagging','expected_loss','existing_control','rencana_mitigasi','kategori_mitigasi']
    d = dict(zip(keys, row)); d.pop('tagging')
    d['urutan'] = len(register) + 1
    register.append(d)

ws = wbv['RAM']
ram = []
for r in range(7, 47):
    code = s(ws.cell(r,4).value)
    ram.append(dict(row=r, kategori_no=s(ws.cell(r,2).value), kode=code, parameter=s(ws.cell(r,5).value), formula=s(ws.cell(r,6).value),
                    satuan=s(ws.cell(r,7).value), kategori_parameter=s(ws.cell(r,8).value),
                    inisiatif=s(ws.cell(r,15).value), kebijakan=s(ws.cell(r,16).value), program=s(ws.cell(r,17).value), anggaran=s(ws.cell(r,18).value)))

wsf = wbf['Rekap']; wsv = wbv['Rekap']
rekap = []
grp = None
for r in range(5, 40):
    b = s(wsv.cell(r,2).value)
    if isinstance(b, str) and not s(wsv.cell(r,3).value): grp = b; continue
    code = s(wsv.cell(r,3).value)
    if not code: continue
    def ref(c):
        f = wsf.cell(r,c).value
        m = re.search(r"\$([A-Z]+)\$(\d+)", str(f)) if f else None
        return f"{m.group(1)}{m.group(2)}" if m else None
    rekap.append(dict(no=int(b), kode=code, parameter=s(wsv.cell(r,4).value), kategori=s(wsv.cell(r,5).value), satuan=s(wsv.cell(r,6).value),
                      grup=grp, diterapkan=s(wsv.cell(r,12).value),
                      ref_limit=ref(7), ref_appetite=ref(8), ref_tolerance=ref(9), ref_hasil=ref(10), ref_status=ref(11)))

ws = wbv['CAT']
cat = []
for r in range(9, 29):
    row = [s(ws.cell(r,c).value) for c in range(3, 16)]
    if any(row):
        keys = ['jenis','nama','rencana_detail','dasar_kewenangan','rencana_waktu','parameter','formula','hasil','rasionalisasi_parameter','limit_v','appetite_v','tolerance_v','rasionalisasi_threshold']
        cat.append(dict(zip(keys, row)))

notes = {sh: V(sh, 'B3') for sh in ['RAS','Kriteria Dampak','Risk Driver','Risk Register','Risiko Utama','RAM','CAT']}
petunjuk = []
ws = wbv['Petunjuk']
for r in range(25, 35):
    petunjuk.append(dict(lembar=s(ws.cell(r,2).value), isi=s(ws.cell(r,4).value)))
notes['Petunjuk'] = V('Petunjuk', 'B3')

json.dump(dict(sampul=sampul, ras=ras, ras_tambahan=ras_tambahan, kategori=kategori, kriteria=kriteria, drivers=drivers,
               register=register, ram=ram, rekap=rekap, cat=cat, notes=notes, petunjuk=petunjuk),
          open(os.path.join(OUT, 'kertas-kerja.json'), 'w'), ensure_ascii=False, indent=0)

# ---------- parameter sheets (cell model) ----------
names = [w.title for w in wbf.worksheets]
params = names[names.index('1a'):]
def color(c):
    try:
        if c is None or c.type != 'rgb': return None
        v = c.rgb
        return v[-6:] if isinstance(v, str) and v not in ('00000000',) else None
    except Exception: return None

RNG = r"\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+"
def patch(f):
    # HyperFormula tidak mendukung kriteria berupa range; ubah ke MMULT setara
    f = re.sub(rf"COUNTIF\(({RNG}),({RNG})\)=0", lambda m: f"MMULT(({m.group(2)}=TRANSPOSE({m.group(1)}))*1,({m.group(1)}={m.group(1)})*1)=0", f)
    f = re.sub(rf"SUMIF\(({RNG}),({RNG}),({RNG})\)", lambda m: f"MMULT(({m.group(2)}=TRANSPOSE({m.group(1)}))*1,IF({m.group(3)}=\"\",0,{m.group(3)}))", f)
    f = re.sub(r'(DATEDIF\([^"]*?)"([a-z]{1,2})"\)', lambda m: f'{m.group(1)}"{m.group(2).upper()}")', f)
    return f

index = []
for t in params:
    wf, wv = wbf[t], wbv[t]
    maxr = maxc = 0
    for row in wf.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != '':
                maxr = max(maxr, c.row); maxc = max(maxc, c.column)
    cells = {}
    for row in wf.iter_rows(min_row=1, max_row=maxr, max_col=maxc):
        for c in row:
            v = c.value
            fill = color(c.fill.fgColor) if c.fill and c.fill.fill_type == 'solid' else None
            if v is None and not fill: continue
            d = {}
            if isinstance(v, ArrayFormula): v = v.text
            if isinstance(v, str) and v.startswith('='):
                d['f'] = patch(v)
                cv = val(wv[c.coordinate].value)
                if cv is not None: d['v'] = cv
            elif v is not None:
                d['v'] = val(v)
            if fill: d['bg'] = fill
            if c.font:
                if c.font.b: d['b'] = 1
                if c.font.i: d['i'] = 1
                fc = color(c.font.color)
                if fc and fc != '000000': d['fc'] = fc
                if c.font.sz and c.font.sz != 9: d['sz'] = c.font.sz
            if c.alignment:
                if c.alignment.horizontal: d['ha'] = c.alignment.horizontal
                if c.alignment.wrap_text: d['w'] = 1
            if c.number_format and c.number_format != 'General': d['nf'] = c.number_format
            if c.border and c.border.left and c.border.left.style: d['bd'] = 1
            cells[c.coordinate] = d
    merges = [str(m) for m in wf.merged_cells.ranges]
    widths = {get_column_letter(i): round(wf.column_dimensions[get_column_letter(i)].width or 8.43, 1) for i in range(1, maxc + 1)}
    dvs = []
    for dv in wf.data_validations.dataValidation:
        if dv.type == 'list' and dv.formula1:
            dvs.append(dict(sqref=str(dv.sqref), list=dv.formula1))
    title = s(wv['E8'].value)
    doc = dict(code=t, title=title, maxRow=maxr, maxCol=maxc, cells=cells, merges=merges, widths=widths, validations=dvs)
    json.dump(doc, open(os.path.join(OUT, 'sheets', f'{t}.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
    index.append(dict(code=t, title=title, maxRow=maxr, maxCol=maxc, ncells=len(cells)))
json.dump(index, open(os.path.join(OUT, 'sheets', 'index.json'), 'w'), ensure_ascii=False, indent=1)
print('ok', len(tax), len(register), len(drivers), len(kriteria), len(rekap), len(index))

# ---------- seed input lembar parameter ----------
kk = json.load(open(os.path.join(OUT, 'kertas-kerja.json')))
pseed = []
for i, rk in enumerate(kk['rekap']):
    d = json.load(open(os.path.join(OUT, 'sheets', f"{rk['kode']}.json")))
    inputs = {a: c['v'] for a, c in d['cells'].items() if c.get('bg') == 'FFFFCC' and 'v' in c and 'f' not in c}
    g = lambda a: d['cells'].get(a, {}).get('v') if a else None
    pseed.append(dict(code=rk['kode'], title=rk['parameter'], grup=rk['grup'], urutan=i + 1, diterapkan=rk['diterapkan'], inputs=inputs,
                      summary=dict(limit=g(rk['ref_limit']), appetite=g(rk['ref_appetite']), tolerance=g(rk['ref_tolerance']),
                                   hasil=g(rk['ref_hasil']), status=g(rk['ref_status']))))
json.dump(pseed, open(os.path.join(OUT, 'parameter-seed.json'), 'w'), ensure_ascii=False)
