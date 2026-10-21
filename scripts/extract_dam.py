# Ekstrak tata letak template DAM (Perencanaan Manajemen Risiko RKAP 2027 PTPN Group) menjadi JSON
# untuk ekspor "Format DAM". Workbook asli terlalu berat dibuka ExcelJS di peramban (external link,
# defined name rusak), jadi sel, gaya, merge, lebar kolom, dan tinggi baris disalin ke JSON statis.
#
# Pemakaian: python3 scripts/extract_dam.py "DAM_Template_Perencanaan Manajemen Risiko RKAP 2027 PTPN Group.xlsx" public/templates/dam
import colorsys, datetime, json, os, re, sys, zipfile
import openpyxl
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
os.makedirs(OUT, exist_ok=True)

SHEETS = {
    "title": "Title", "ras-cover": "RAS==>", "ras": "Risk.RAS", "ram-cover": "RAM==>", "ram": "Risk.RAM-form isian",
    "register": "Risk Register", "driver": "Risk Driver", "kriteria": "Kriteria Dampak",
    "standar": "Risk.RAM-Parameter Standar", "klaster": "Risk RAM - Parameter Klaster",
    "cat-cover": "CAT==>", "cat": "Risk.CAT-aksikorporasi",
}

# ---------- warna tema (openpyxl tidak menerjemahkan theme + tint) ----------
with zipfile.ZipFile(SRC) as z:
    theme = z.read("xl/theme/theme1.xml").decode()
    for name in ("image1.png", "image2.png"):
        open(os.path.join(OUT, {"image1.png": "logo.png", "image2.png": "framework.png"}[name]), "wb").write(z.read(f"xl/media/{name}"))
scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", theme, re.S).group(0)
order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
THEME = []
for k in order:
    m = re.search(rf"<a:{k}>.*?(?:srgbClr val=\"([0-9A-Fa-f]{{6}})\"|lastClr=\"([0-9A-Fa-f]{{6}})\").*?</a:{k}>", scheme, re.S)
    THEME.append((m.group(1) or m.group(2)).upper())

def tinted(hexrgb, tint):
    if not tint: return hexrgb
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    return "".join(f"{round(x * 255):02X}" for x in colorsys.hls_to_rgb(h, l, s))

def color(c):
    try:
        if c is None: return None
        if c.type == "rgb" and isinstance(c.rgb, str) and c.rgb not in ("00000000",): return c.rgb[-6:].upper()
        if c.type == "theme" and c.theme is not None: return tinted(THEME[c.theme], c.tint)
    except Exception: return None
    return None

EPOCH = datetime.datetime(1899, 12, 30)
def val(v):
    if isinstance(v, datetime.datetime): return round((v - EPOCH).total_seconds() / 86400, 6)
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v

wb = openpyxl.load_workbook(SRC)
index = {}
for slug, name in SHEETS.items():
    ws = wb[name]
    maxr = maxc = 0
    for row in ws.iter_rows():
        for c in row:
            has_fill = c.fill and c.fill.fill_type == "solid" and color(c.fill.fgColor) not in (None, "FFFFFF")
            if (c.value is not None and str(c.value).strip() != "") or has_fill or (c.border and c.border.left and c.border.left.style):
                maxr, maxc = max(maxr, c.row), max(maxc, c.column)
    cells = {}
    for row in ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc):
        for c in row:
            v = c.value
            fill = color(c.fill.fgColor) if c.fill and c.fill.fill_type == "solid" else None
            bd = c.border and any(getattr(c.border, s) is not None and getattr(c.border, s).style for s in ("left", "right", "top", "bottom"))
            if (v is None or str(v).strip() == "") and not fill and not bd: continue
            d = {}
            if isinstance(v, str) and v.startswith("="): d["f"] = v
            elif v is not None and str(v).strip() != "": d["v"] = val(v)
            if fill: d["bg"] = fill
            if c.font:
                if c.font.b: d["b"] = 1
                if c.font.i: d["i"] = 1
                fc = color(c.font.color)
                if fc and fc != "000000": d["fc"] = fc
                if c.font.sz: d["sz"] = c.font.sz
                if c.font.name and c.font.name != "Arial": d["fn"] = c.font.name
            if c.alignment:
                if c.alignment.horizontal: d["ha"] = c.alignment.horizontal
                if c.alignment.vertical: d["va"] = c.alignment.vertical
                if c.alignment.wrap_text: d["w"] = 1
            if c.number_format and c.number_format != "General": d["nf"] = c.number_format
            if bd: d["bd"] = 1
            cells[c.coordinate] = d
    # Lembar sampul berlatar penuh (1000 baris): gaya dominan disimpan sekali sebagai "base",
    # sel yang sama persis dengan base dibuang, dan area latar dibatasi 60 baris.
    base = None
    sig = lambda c: json.dumps({k: v for k, v in c.items() if k != "v"}, sort_keys=True)
    counts = {}
    for c in cells.values(): counts[sig(c)] = counts.get(sig(c), 0) + 1
    top = max(counts, key=counts.get) if cells else None
    if top and max(counts.values()) > 0.5 * maxr * maxc and json.loads(top).get("bg") not in (None, "FFFFFF"):
        base = json.loads(top)
        cells = {a: c for a, c in cells.items() if "v" in c or sig(c) != top}
        maxr = min(maxr, 60)
    widths = {get_column_letter(i): round(ws.column_dimensions[get_column_letter(i)].width or 8.43, 1) for i in range(1, maxc + 1)}
    heights = {r: ws.row_dimensions[r].height for r in range(1, maxr + 1) if ws.row_dimensions[r].height}
    doc = dict(code=slug, title=name, maxRow=maxr, maxCol=maxc, cells=cells, merges=[str(m) for m in ws.merged_cells.ranges],
               widths=widths, heights=heights, validations=[], tab=color(ws.sheet_properties.tabColor), grid=ws.sheet_view.showGridLines, base=base)
    json.dump(doc, open(os.path.join(OUT, f"{slug}.json"), "w"), ensure_ascii=False, separators=(",", ":"))
    index[slug] = dict(name=name, maxRow=maxr, maxCol=maxc, ncells=len(cells))
print(json.dumps(index, indent=1))
